"""Genera el informe Word desde Markdown usando python-docx del runtime de documentos."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'
DIAG=OUT/'diagramas';DIAG.mkdir(exist_ok=True)
FONT='C:/Windows/Fonts/arial.ttf'
BOLD='C:/Windows/Fonts/arialbd.ttf'
def font(size,bold=False):return ImageFont.truetype(BOLD if bold else FONT,size)
def box(draw,rect,title,lines):
    x,y,x2,y2=rect
    draw.rounded_rectangle(rect,radius=14,fill='#F4F6F9',outline='#A7B9CD',width=2)
    draw.text((x+22,y+18),title,font=font(26,True),fill='#234E76')
    for n,line in enumerate(lines):draw.text((x+22,y+62+n*32),line,font=font(21),fill='#344454')
def arrow(draw,a,b,label=None):
    draw.line([a,b],fill='#547899',width=4)
    x,y=b
    if b[0]>a[0]:draw.polygon([(x,y),(x-13,y-8),(x-13,y+8)],fill='#547899')
    else:draw.polygon([(x,y),(x-8,y-13),(x+8,y-13)],fill='#547899')
    if label:draw.text(((a[0]+b[0])/2-30,(a[1]+b[1])/2-28),label,font=font(18),fill='#547899')
im=Image.new('RGB',(1400,390),'white');dr=ImageDraw.Draw(im)
box(dr,(15,55,300,230),'Navegador',['HTML5 / CSS3','Formularios GET / POST'])
box(dr,(380,55,690,230),'Tomcat 11',['HotelServlet + Rules','JSP: presentación'])
box(dr,(770,55,1070,230),'Persistencia Java',['HotelDao','Database + JDBC'])
box(dr,(1140,55,1390,230),'MySQL',['hotel_reservas','InnoDB'])
arrow(dr,(300,140),(380,140));arrow(dr,(690,140),(770,140));arrow(dr,(1070,140),(1140,140))
dr.text((25,295),'La JSP devuelve HTML. El DAO valida disponibilidad dentro de una transacción.',font=font(24),fill='#566779')
im.save(DIAG/'arquitectura.png')
im=Image.new('RGB',(1400,880),'white');dr=ImageDraw.Draw(im)
box(dr,(20,25,420,325),'CLIENTES',['PK id · UK documento','nombre, telefono, email','fecha_registro'])
box(dr,(930,25,1380,325),'HABITACIONES',['PK id · UK numero','tipo, capacidad, habilitada','descripcion, precio_noche'])
box(dr,(470,380,930,850),'RESERVAS',['PK id','FK cliente_id','FK habitacion_id','fecha_entrada, fecha_salida','noches, numero_huespedes','precio_noche, total','estado, fecha_creacion','observaciones'])
arrow(dr,(420,225),(570,380),'1 → N');arrow(dr,(930,225),(830,380),'1 → N')
box(dr,(20,470,400,770),'USUARIOS',['PK id · UK username','nombre, password, rol','activo','Sin uso en este avance'])
im.save(DIAG/'modelo-er.png')

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492)
sec.different_first_page_header_footer=True
styles=doc.styles
for name,size,color,before,after in [('Normal',11,'202020',0,6),('Title',25,'173A57',0,16),('Subtitle',14,'5B6670',0,10),('Heading 1',16,'2E74B5',16,8),('Heading 2',13,'2E74B5',12,6),('Heading 3',12,'1F4D78',8,4)]:
    s=styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after);s.paragraph_format.line_spacing=1.10
    if name.startswith('Heading'):s.paragraph_format.keep_with_next=True
for name,size,color in [('Code',8.5,'24415B'),('Caption',9,'6D747B'),('CoverMeta',11,'303030'),('IndexDetail',10,'505050')]:
    s=styles[name] if name in styles else styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.font.name='Consolas' if name=='Code' else 'Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(4);s.paragraph_format.line_spacing=1.10
styles['CoverMeta'].paragraph_format.space_after=Pt(12)
styles['IndexDetail'].paragraph_format.left_indent=Inches(.25)
header=sec.header.paragraphs[0];header.text='DESARROLLO WEB INTEGRADO  |  APF1';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.style='Caption';footer.add_run('Sistema de Gestión y Reservas Hoteleras  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

numbering=doc.part.numbering_part.element
num_counter=80
def number_list(p,start=1,bullet=False):
    global num_counter
    num_counter+=1;aid=str(num_counter)
    abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),aid)
    lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
    for tag,value in [('start',str(start)),('numFmt','bullet' if bullet else 'decimal'),('lvlText','•' if bullet else '%1.'),('lvlJc','left')]:
        el=OxmlElement('w:'+tag);el.set(qn('w:val'),value);lvl.append(el)
    pp=OxmlElement('w:pPr');indent=OxmlElement('w:ind');indent.set(qn('w:left'),'720');indent.set(qn('w:hanging'),'360');pp.append(indent)
    tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'num');tab.set(qn('w:pos'),'720');tabs.append(tab);pp.append(tabs);lvl.append(pp);abstract.append(lvl);numbering.append(abstract)
    num=OxmlElement('w:num');num.set(qn('w:numId'),aid);abstract_id=OxmlElement('w:abstractNumId');abstract_id.set(qn('w:val'),aid);num.append(abstract_id);numbering.append(num)
    props=p._p.get_or_add_pPr();np=OxmlElement('w:numPr');level=OxmlElement('w:ilvl');level.set(qn('w:val'),'0');ni=OxmlElement('w:numId');ni.set(qn('w:val'),aid);np.extend([level,ni]);props.append(np)
    p.paragraph_format.line_spacing=1.167;p.paragraph_format.space_after=Pt(8)
def table(lines):
    rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in lines if not re.match(r'^\|[\s:|\-]+\|$',line)]
    tab=doc.add_table(rows=0,cols=len(rows[0]));tab.autofit=False
    widths=[3000,6360] if len(rows[0])==2 else [9360//len(rows[0])]*len(rows[0])
    pr=tab._tbl.tblPr
    tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'9360')
    ind=OxmlElement('w:tblInd');ind.set(qn('w:w'),'120');ind.set(qn('w:type'),'dxa');pr.append(ind)
    margins=OxmlElement('w:tblCellMar')
    for k,v in [('top',80),('bottom',80),('start',120),('end',120)]:
        n=OxmlElement('w:'+k);n.set(qn('w:w'),str(v));n.set(qn('w:type'),'dxa');margins.append(n)
    pr.append(margins)
    borders=OxmlElement('w:tblBorders')
    for k in ['top','left','bottom','right','insideH','insideV']:
        n=OxmlElement('w:'+k);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');n.set(qn('w:color'),'D5DCE2');borders.append(n)
    pr.append(borders)
    for col,width in zip(tab._tbl.tblGrid.gridCol_lst,widths):col.set(qn('w:w'),str(width))
    for i,row in enumerate(rows):
        cells=tab.add_row().cells
        for cell,text,width in zip(cells,row,widths):
            cell.width=Inches(width/1440);cell._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(width));cell.text=text
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
            for r in p.runs:r.font.size=Pt(9);r.bold=(i==0)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F4F7');cell._tc.get_or_add_tcPr().append(sh)
        if i==0:
            trpr=tab.rows[-1]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=(OUT/'Informe_APF1.md').read_text(encoding='utf-8').splitlines();i=0;page=0;code=False
while i<len(lines):
    line=lines[i];i+=1
    if line=='<!-- PAGEBREAK -->':doc.add_page_break();page+=1;continue
    if line.startswith('```'):code=not code;continue
    if code:
        p=doc.add_paragraph(line,style='Code');p.paragraph_format.keep_with_next=i<len(lines) and not lines[i].startswith('```');continue
    if not line.strip():continue
    if line.startswith('|'):
        group=[line]
        while i<len(lines) and lines[i].startswith('|'):group.append(lines[i]);i+=1
        table(group);continue
    img=re.match(r'!\[(.*?)\]\((.*?)\)',line)
    if img:
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(OUT/img[2]),width=Inches(6.5))
        doc.add_paragraph(img[1],style='Caption');continue
    if page==0:
        if line.startswith('# '):
            p=doc.add_paragraph(line[2:],style='Title');p.paragraph_format.space_before=Pt(40)
        elif line=='Sistema de Gestión y Reservas Hoteleras':
            p=doc.add_paragraph(line,style='Title');p.paragraph_format.space_before=Pt(22);p.paragraph_format.space_after=Pt(26)
        elif re.match(r'^\d\. ',line):
            p=doc.add_paragraph(line[3:],style='CoverMeta');number_list(p,int(line[0]));continue
        else:p=doc.add_paragraph(line,style='CoverMeta')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;continue
    if line.startswith('## '):doc.add_heading(line[3:],2);continue
    if line.startswith('# '):doc.add_heading(line[2:],1);continue
    ordered=re.match(r'^(\d+)\. (.*)',line)
    if ordered:
        p=doc.add_paragraph(ordered[2]);number_list(p,int(ordered[1]));continue
    if line.startswith('- '):
        p=doc.add_paragraph(line[2:]);number_list(p,bullet=True);continue
    if line.startswith('   '):doc.add_paragraph(line.strip(),style='IndexDetail');continue
    doc.add_paragraph(line)
doc.core_properties.title='Sistema de Gestión y Reservas Hoteleras — APF1'
doc.core_properties.author='Equipo APF1'
doc.core_properties.subject='Desarrollo Web Integrado'
report_path=OUT/'Informe_APF1.docx'
try:
    doc.save(report_path)
except PermissionError:
    report_path=OUT/'Informe_APF1_actualizado.docx'
    doc.save(report_path)
# Comprobación estructural: encabezados y tablas, sin atribuir revisión visual.
check=Document(report_path)
assert len([p for p in check.paragraphs if p.style.name=='Heading 1'])==8
assert len(check.tables)==5
assert all(sum(int(c.get(qn('w:w'))) for c in t._tbl.tblGrid.gridCol_lst)==9360 for t in check.tables)
(OUT/'evidencias'/'documento-qa.txt').write_text('Informe generado con python-docx. Preset standard_business_brief; portada editorial adaptada al enunciado.\nGeometría Letter, márgenes 1 pulgada; tablas DXA 9360, estilos y numeración explícitos.\nComprobación estructural correcta. Revisión visual pendiente: LibreOffice no disponible.\n',encoding='utf-8')
print('Informe Word generado; estructura comprobada.')
